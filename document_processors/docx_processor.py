"""
document_processors/docx_processor.py
Extract structured text from Word documents.
"""
from docx import Document


def extract_text(file_path: str) -> str:
    doc = Document(file_path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    # Tables
    for i, table in enumerate(doc.tables, 1):
        lines.append(f"\n[Table {i}]")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)
