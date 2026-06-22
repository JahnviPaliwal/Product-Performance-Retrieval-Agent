"""
document_processors/pptx_processor.py
Extract text from PowerPoint slides and speaker notes.
"""
from pptx import Presentation


def extract_text(file_path: str) -> str:
    prs = Presentation(file_path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        parts = [f"[Slide {i}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
        # Notes
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notes: {notes}")
        slides.append("\n".join(parts))
    return "\n\n".join(slides)
